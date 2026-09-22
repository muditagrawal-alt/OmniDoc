"""
Docling Document Parser for OmniDoc.
Converts PDF and DOCX into structured Markdown ASTs, extracting tables,
figure bounding boxes, and hierarchical section-aware chunks.
"""
import os
import re
import uuid
import logging
from typing import List, Dict, Any, Optional
from dataclasses import dataclass

from core.state import RetrievedChunk, VisualElement

logger = logging.getLogger("OmniDoc.DoclingParser")


@dataclass
class ParsedDocument:
    doc_id: str
    filename: str
    markdown_content: str
    chunks: List[RetrievedChunk]
    visual_elements: List[VisualElement]
    tables: List[Dict[str, Any]]


class DoclingParser:
    """Parses documents with layout awareness using IBM Docling."""

    def __init__(self, extracted_images_dir: str = "extracted_images"):
        self.images_dir = extracted_images_dir
        os.makedirs(self.images_dir, exist_ok=True)
        self.converter = None
        self._init_converter()

    def _init_converter(self):
        try:
            from docling.document_converter import DocumentConverter
            self.converter = DocumentConverter()
            logger.info("Docling DocumentConverter ready.")
        except Exception as e:
            logger.warning(f"Docling initialization notice: {e}. PyMuPDF fallback available.")

    def parse_document(self, file_path: str, doc_id: Optional[str] = None) -> ParsedDocument:
        """
        Parses a PDF or DOCX file into structured text, tables, and figures.
        """
        filename = os.path.basename(file_path)
        doc_id = doc_id or f"doc_{uuid.uuid4().hex[:10]}"

        if self.converter:
            try:
                return self._parse_with_docling(file_path, doc_id, filename)
            except Exception as e:
                logger.error(f"Docling parsing failed on {filename} ({e}). Using PyMuPDF fallback.")

        return self._parse_with_pymupdf_fallback(file_path, doc_id, filename)

    def _parse_with_docling(self, file_path: str, doc_id: str, filename: str) -> ParsedDocument:
        """Primary Docling conversion path with HybridChunker."""
        from docling.chunking import HybridChunker

        conv_res = self.converter.convert(file_path)
        doc = conv_res.document
        markdown_text = doc.export_to_markdown()

        # Extract hierarchical chunks
        chunker = HybridChunker(max_tokens=500)
        chunk_iter = chunker.chunk(doc)
        
        chunks: List[RetrievedChunk] = []
        for i, ch in enumerate(chunk_iter, 1):
            text = ch.text.strip()
            if not text:
                continue
            
            # Extract section heading if available in metadata
            section = "General"
            if hasattr(ch, "meta") and ch.meta:
                headings = getattr(ch.meta, "headings", [])
                if headings:
                    section = " > ".join(headings)
                    
            page_num = 1
            if hasattr(ch, "meta") and hasattr(ch.meta, "doc_items"):
                items = ch.meta.doc_items
                if items and hasattr(items[0], "prov") and items[0].prov:
                    page_num = items[0].prov[0].page_no

            chunks.append(RetrievedChunk(
                chunk_id=f"{doc_id}_c{i:03d}",
                doc_id=doc_id,
                text=text,
                page_number=page_num,
                section_title=section,
                score=0.0,
                retrieval_method="hybrid"
            ))

        # Extract tables
        tables = []
        for i, tbl in enumerate(getattr(doc, "tables", []), 1):
            try:
                tbl_md = tbl.export_to_markdown()
                tables.append({
                    "table_id": f"{doc_id}_tbl_{i}",
                    "markdown": tbl_md,
                    "page_number": getattr(tbl, "page_no", 1)
                })
            except Exception:
                pass

        # Extract figures
        visual_elements: List[VisualElement] = []
        for i, fig in enumerate(getattr(doc, "pictures", []), 1):
            fig_id = f"{doc_id}_fig_{i}"
            caption = getattr(fig, "caption", "") or f"Figure {i} in {filename}"
            bbox = None
            if hasattr(fig, "prov") and fig.prov:
                b = fig.prov[0].bbox
                bbox = [b.l, b.t, b.r, b.b] if b else None

            visual_elements.append(VisualElement(
                figure_id=fig_id,
                doc_id=doc_id,
                page_number=getattr(fig, "page_no", 1),
                caption=caption,
                image_path=None,
                bounding_box=bbox
            ))

        return ParsedDocument(
            doc_id=doc_id,
            filename=filename,
            markdown_content=markdown_text,
            chunks=chunks,
            visual_elements=visual_elements,
            tables=tables
        )

    def _parse_with_pymupdf_fallback(self, file_path: str, doc_id: str, filename: str) -> ParsedDocument:
        """Robust fallback using PyMuPDF (fitz) or python-docx."""
        chunks: List[RetrievedChunk] = []
        full_text_parts = []
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            import fitz
            doc = fitz.open(file_path)
            for page_idx, page in enumerate(doc, 1):
                p_text = page.get_text()
                if p_text.strip():
                    full_text_parts.append(p_text)
                    # Window chunks per page
                    words = p_text.split()
                    chunk_size = 400
                    for c_idx in range(0, len(words), chunk_size - 80):
                        chunk_words = words[c_idx:c_idx + chunk_size]
                        chunks.append(RetrievedChunk(
                            chunk_id=f"{doc_id}_p{page_idx}_c{c_idx}",
                            doc_id=doc_id,
                            text=" ".join(chunk_words),
                            page_number=page_idx,
                            section_title=f"Page {page_idx}",
                            score=0.0
                        ))
        else:
            import docx
            doc = docx.Document(file_path)
            paras = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text_parts = paras
            text_blob = "\n\n".join(paras)
            words = text_blob.split()
            for c_idx in range(0, len(words), 350):
                chunks.append(RetrievedChunk(
                    chunk_id=f"{doc_id}_c{c_idx}",
                    doc_id=doc_id,
                    text=" ".join(words[c_idx:c_idx + 400]),
                    page_number=1,
                    section_title="General",
                    score=0.0
                ))

        return ParsedDocument(
            doc_id=doc_id,
            filename=filename,
            markdown_content="\n\n".join(full_text_parts),
            chunks=chunks,
            visual_elements=[],
            tables=[]
        )
