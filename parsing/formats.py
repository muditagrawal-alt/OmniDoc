"""
More document formats for the fast parser: PowerPoint slides, e-mail (with attachments),
EPUB books, saved web pages, and audio/video recordings (transcribed with timestamps).

Every reader returns blocks ``(page, section, text)`` like the other formats: a slide or
chapter number as the page, the slide title / chapter / timestamp as the section.
"""
import os
import re
import io
import email
import zipfile
import logging
from email import policy
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger("OmniDoc.Formats")

Block = Tuple[int, str, str]
MEDIA_EXTENSIONS = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".webm", ".mp4", ".mov", ".mpeg", ".mpga", ".aac"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".webm", ".mpeg"}
SEGMENT_SECONDS = 45  # transcript lines are grouped into blocks of about this length


# --------------------------------------------------------------------------- slides
def pptx_blocks(path: str) -> Tuple[List[Block], List[Dict[str, Any]]]:
    """Slides as pages (title as section), text frames in reading order, tables, and speaker notes."""
    from pptx import Presentation
    prs = Presentation(path)
    blocks: List[Block] = []
    tables: List[Dict[str, Any]] = []
    for number, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        section = (title or f"Slide {number}")[:120]
        shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
        for shape in shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                if text:
                    blocks.append((number, section, text))
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                if len(rows) >= 2:
                    tables.append({"page": number, "bbox": None, "title": f"{section} (table)", "columns": rows[0],
                                   "rows": rows[1:], "source": "pptx"})
                    blocks.append((number, section, "\n".join(" | ".join(r) for r in rows)))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                blocks.append((number, section, f"Speaker notes: {notes}"))
        if title and not any(b[0] == number for b in blocks):
            blocks.append((number, section, title))
    return blocks, tables


# ---------------------------------------------------------------------------- email
def _html_text(html: str) -> List[str]:
    from parsing.web import readable
    return readable(html).get("paragraphs") or []


def eml_blocks(path: str) -> Tuple[List[Block], List[str]]:
    """The message headers, the body, and the text of readable attachments. Returns (blocks, attachment names)."""
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
    header = "\n".join(f"{h}: {msg[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if msg[h])
    blocks: List[Block] = [(1, "Message", header)] if header else []
    subject = str(msg["Subject"] or "Message")[:120]
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is not None:
        content = body.get_content()
        paragraphs = _html_text(content) if body.get_content_type() == "text/html" else \
            [p.strip() for p in re.split(r"\n\s*\n", content) if p.strip()]
        blocks.extend((1, subject, p) for p in paragraphs)
    names: List[str] = []
    for part in msg.iter_attachments():
        name = part.get_filename() or "attachment"
        names.append(name)
        data = part.get_payload(decode=True) or b""
        text = _attachment_text(name, data)
        if text:
            blocks.extend((1, f"Attachment: {name}"[:120], p) for p in text)
    return blocks, names


def _attachment_text(name: str, data: bytes) -> List[str]:
    ext = os.path.splitext(name)[1].lower()
    try:
        if ext == ".pdf":
            import fitz
            with fitz.open("pdf", data) as doc:
                text = "\n\n".join(page.get_text() for page in doc[:50])
        elif ext == ".docx":
            import docx
            text = "\n\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs if p.text.strip())
        elif ext in (".txt", ".md", ".csv"):
            text = data.decode("utf-8", errors="replace")
        elif ext in (".html", ".htm"):
            return _html_text(data.decode("utf-8", errors="replace"))
        else:
            return []
    except Exception as e:
        logger.info(f"Attachment {name} skipped: {e}")
        return []
    return [re.sub(r"[ \t]+", " ", p).strip() for p in re.split(r"\n\s*\n", text) if p.strip()][:500]


# ----------------------------------------------------------------------------- epub
def epub_blocks(path: str) -> List[Block]:
    """Chapters in reading order (spine) as pages, their headings as sections."""
    from bs4 import BeautifulSoup
    blocks: List[Block] = []
    with zipfile.ZipFile(path) as z:
        container = BeautifulSoup(z.read("META-INF/container.xml"), "xml")
        opf_path = container.find("rootfile")["full-path"]
        opf = BeautifulSoup(z.read(opf_path), "xml")
        base = os.path.dirname(opf_path)
        manifest = {item["id"]: item["href"] for item in opf.find_all("item") if item.get("id") and item.get("href")}
        order = [manifest[ref["idref"]] for ref in opf.find_all("itemref") if ref.get("idref") in manifest]
        for chapter, href in enumerate(order, 1):
            name = os.path.normpath(os.path.join(base, href)).replace("\\", "/")
            try:
                html = z.read(name).decode("utf-8", errors="replace")
            except KeyError:
                continue
            section = f"Chapter {chapter}"
            for para in _html_text(html):
                if para.startswith("## "):
                    section = para[3:][:120]
                    continue
                blocks.append((chapter, section, para))
    return blocks


def html_blocks(path: str) -> List[Block]:
    with open(path, "rb") as f:
        html = f.read().decode("utf-8", errors="replace")
    blocks: List[Block] = []
    section = ""
    for para in _html_text(html):
        if para.startswith("## "):
            section = para[3:][:120]
            continue
        blocks.append((1, section, para))
    return blocks


# ---------------------------------------------------------------------------- media
def stamp(seconds: float) -> str:
    seconds = int(seconds or 0)
    h, rem = divmod(seconds, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def media_blocks(path: str) -> Tuple[List[Block], Dict[str, Any]]:
    """
    Transcribes a recording and groups the transcript into blocks of about SEGMENT_SECONDS,
    each with its time range as the section. Returns (blocks, media info with the segments).
    """
    from parsing.transcribe import transcribe
    result = transcribe(path)
    segments = result["segments"]
    blocks: List[Block] = []
    group: List[Dict[str, Any]] = []

    def flush() -> None:
        if group:
            start, end = group[0]["start"], group[-1]["end"]
            blocks.append((1, f"{stamp(start)} – {stamp(end)}", " ".join(s["text"].strip() for s in group)))
            group.clear()

    for seg in segments:
        if group and seg["end"] - group[0]["start"] > SEGMENT_SECONDS:
            flush()
        group.append(seg)
    flush()
    ext = os.path.splitext(path)[1].lower()
    info = {"kind": "video" if ext in VIDEO_EXTENSIONS else "audio", "language": result.get("language"),
            "duration": result.get("duration"), "engine": result.get("engine"),
            "segments": [{"start": round(s["start"], 2), "end": round(s["end"], 2), "text": s["text"].strip()} for s in segments]}
    return blocks, info


def section_start(section: str) -> Optional[float]:
    """Seconds at the start of a "mm:ss – mm:ss" section (None if it is not a time range)."""
    m = re.match(r"^(?:(\d+):)?(\d{1,2}):(\d{2})\s*[–-]", section or "")
    if not m:
        return None
    h, mi, s = int(m.group(1) or 0), int(m.group(2)), int(m.group(3))
    return float(h * 3600 + mi * 60 + s)
